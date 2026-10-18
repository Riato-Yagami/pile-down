#!/usr/bin/env python3
"""Build the formatted, illustrated workbook from the archived prompt history.

The CSV remains the editable source of summaries. The Markdown identifies the
original session records so images can be recovered without temporary files.
"""

from __future__ import annotations

import argparse
import base64
import csv
from datetime import datetime
from html import unescape
from io import BytesIO
import json
import math
from pathlib import Path
import re
import sys

try:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as WorkbookImage
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
    from openpyxl.drawing.xdr import XDRPositiveSize2D
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.units import pixels_to_EMU
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from PIL import Image, ImageOps
except ImportError as error:
    raise SystemExit(
        "Dépendance manquante. Exécuter : python -m pip install -r "
        "wip/prompt/requirements.txt"
    ) from error


KNOWN_COLUMNS = {
    "id", "date", "prompt", "summary", "images", "active_file", "open_tabs",
    "files", "selection", "question", "source",
}
IMAGE_NOTE = re.compile(r"\s*\[Images jointes : \d+ ; fichiers image non inclus dans le CSV\.\]\s*$")
REQUEST = re.compile(r"^## My request(?: for Codex)?:[ \t]*\n?", re.MULTILINE)
IDE_SECTION = re.compile(
    r"^## (Active file|Active selection of the file|Open tabs):[ \t]*(.*)$",
    re.MULTILINE,
)
NO_SUMMARY = "Aucun bilan final distinct"


def path_from_config(value: str, directory: Path) -> Path:
    path = Path(value).expanduser()
    return path if path.is_absolute() else directory / path


def parse_prompt(raw: str, read_pasted_files: bool) -> dict:
    """Separate IDE scaffolding from the actual request, without dropping context."""
    text = unescape(IMAGE_NOTE.sub("", raw)).strip()
    values = dict.fromkeys(["active_file", "open_tabs", "selection", "files", "question"], "")
    if text.startswith("<send_user_message_question_reply>"):
        data = text.removeprefix("<send_user_message_question_reply>")
        data = data.removesuffix("</send_user_message_question_reply>").strip()
        replies = json.loads(data)
        values["question"] = "\n\n".join(item.get("question", "") for item in replies)
        values["prompt"] = "\n\n".join(item.get("answer", "") for item in replies)
        return values

    marker = REQUEST.search(text)
    context = text[:marker.start()] if marker else ""
    request = text[marker.end():].strip() if marker else text
    headings = list(IDE_SECTION.finditer(context))
    keys = {"Active file": "active_file", "Open tabs": "open_tabs", "Active selection of the file": "selection"}
    for index, heading in enumerate(headings):
        end = headings[index + 1].start() if index + 1 < len(headings) else len(context)
        value = (heading.group(2) + "\n" + context[heading.end():end]).strip()
        if heading.group(1) == "Open tabs":
            value = "\n".join(line.removeprefix("- ") for line in value.splitlines())
        values[keys[heading.group(1)]] = value

    attached_files = re.findall(r"^## .*?: ([A-Za-z]:[\\/].+)$", context, re.MULTILINE)
    values["files"] = "\n".join(attached_files)
    if not request and attached_files and read_pasted_files:
        pasted = []
        for filename in attached_files:
            path = Path(filename.strip())
            if path.suffix.lower() in {".txt", ".md"} and path.is_file():
                pasted.append(path.read_text(encoding="utf-8-sig"))
        request = "\n\n".join(pasted).strip()
    if not request and values["selection"]:
        request = "Demande transmise dans la sélection du fichier : voir l'onglet Contexte."
    if not request and attached_files:
        request = "Demande transmise dans un fichier joint : voir l'onglet Contexte."
    # The actual screenshots are embedded rather than represented by Markdown tags.
    request = re.sub(r"\\?!\[[^\]]*\]\\?\([^)]*\)", "", request).strip()
    values["prompt"] = request or "Image jointe."
    return values


def load_records(config: dict, directory: Path) -> list[dict]:
    csv.field_size_limit(10_000_000)
    with path_from_config(config["input_csv"], directory).open(encoding="utf-8-sig", newline="") as stream:
        source_rows = list(csv.DictReader(stream, delimiter=";"))
    archive = path_from_config(config["input_archive"], directory).read_text(encoding="utf-8")
    sources = {}
    for block in re.split(r"(?=^### Message \d+ - )", archive, flags=re.MULTILINE)[1:]:
        number = re.match(r"### Message (\d+) - ", block)
        source = re.search(r"^Source : `\.codex/(.*?)`, ligne (\d+)\.", block, re.MULTILINE)
        if number and source:
            sources[int(number.group(1))] = (source.group(1), int(source.group(2)))

    filters = config["filters"]
    records = []
    for number, row in enumerate(source_rows, 1):
        if number in filters["exclude_ids"]:
            continue
        if filters["include_ids"] and number not in filters["include_ids"]:
            continue
        date = datetime.fromisoformat(row["date"])
        if filters["date_from"] and date.date().isoformat() < filters["date_from"]:
            continue
        if filters["date_to"] and date.date().isoformat() > filters["date_to"]:
            continue
        summary = row["résumé des changements"]
        if filters["only_with_summary"] and summary.startswith(NO_SUMMARY):
            continue
        if filters["contains"] and not any(word.casefold() in (row["prompt"] + summary).casefold() for word in filters["contains"]):
            continue
        fields = parse_prompt(row["prompt"], config["text"]["read_pasted_files"])
        note = re.search(r"\[Images jointes : (\d+) ;", row["prompt"])
        fields.update(
            id=number, date=date.replace(tzinfo=None), summary=summary,
            source=sources.get(number), image_data=[], expected_images=int(note.group(1)) if note else 0,
        )
        records.append(fields)
    records.sort(key=lambda row: (row["date"], row["id"]), reverse=filters["order"] == "descending")
    if not records:
        raise ValueError("Les filtres ne sélectionnent aucun prompt.")
    return records


def recover_images(records: list[dict], config: dict, directory: Path) -> list[str]:
    if not config["images"]["enabled"]:
        return []
    codex_home = path_from_config(config["codex_home"], directory).resolve()
    grouped = {}
    warnings = []
    for record in records:
        if record["expected_images"] and record["source"]:
            relative, line = record["source"]
            grouped.setdefault(relative, {})[line] = record
    for relative, wanted in grouped.items():
        path = (codex_home / relative).resolve()
        if not path.is_relative_to(codex_home):
            raise ValueError(f"Source hors du dossier Codex : {relative}")
        if not path.is_file():
            warnings.append(f"Session introuvable : {relative}")
            continue
        last_line = max(wanted)
        with path.open(encoding="utf-8") as stream:
            for number, line in enumerate(stream, 1):
                if number in wanted:
                    payload = json.loads(line).get("payload", {})
                    if payload.get("role") != "user":
                        raise ValueError(f"La source du prompt {wanted[number]['id']} n'est plus un message utilisateur.")
                    for part in payload.get("content", []):
                        if part.get("type") not in {"input_image", "image"}:
                            continue
                        url = part.get("image_url", part.get("url", ""))
                        if isinstance(url, dict):
                            url = url.get("url", "")
                        if url.startswith("data:image/") and ";base64," in url:
                            wanted[number]["image_data"].append(base64.b64decode(url.split(",", 1)[1], validate=True))
                        else:
                            warnings.append(f"Prompt {wanted[number]['id']} : format d'image non exportable automatiquement.")
                if number >= last_line:
                    break
    incomplete = [str(r["id"]) for r in records if len(r["image_data"]) != r["expected_images"]]
    if incomplete:
        message = "Images manquantes pour les prompts : " + ", ".join(incomplete)
        if config["images"]["fail_on_missing"]:
            raise ValueError(message + ". Vérifier codex_home, ou désactiver images.fail_on_missing.")
        warnings.append(message)
    return warnings


def add_picture(sheet, content: bytes, column: int, row: int, width: int, height: int, top: int) -> int:
    with Image.open(BytesIO(content)) as original:
        image = ImageOps.exif_transpose(original).convert("RGBA")
        image.thumbnail((width, height), Image.Resampling.LANCZOS)
        buffer = BytesIO()
        image.save(buffer, format="PNG", optimize=True)
        buffer.seek(0)
        picture = WorkbookImage(buffer)
        picture.anchor = OneCellAnchor(
            _from=AnchorMarker(col=column - 1, row=row - 1, colOff=pixels_to_EMU(6), rowOff=pixels_to_EMU(top)),
            ext=XDRPositiveSize2D(pixels_to_EMU(image.width), pixels_to_EMU(image.height)),
        )
        sheet.add_image(picture)
        return image.height


def format_table(sheet, columns: list[dict], name: str, style: dict) -> None:
    sheet.freeze_panes = "C2" if sheet.title == "Historique" else "B2"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = style["zoom"]
    sheet.row_dimensions[1].height = 30
    for index, column in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(index)].width = column["width"]
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = Font(name=style["font"], size=style["font_size"])
            if isinstance(cell.value, str):
                # Requests are text, including those that happen to start with '='.
                cell.data_type = "s"
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor=style["header_background"])
                cell.font = Font(name=style["font"], size=style["font_size"], bold=True, color=style["header_text"])
            elif cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=style["stripe_background"])
            if isinstance(cell.value, datetime):
                cell.number_format = "dd/mm/yyyy hh:mm"
    table = Table(displayName=name, ref=f"A1:{get_column_letter(len(columns))}{sheet.max_row}")
    table.tableStyleInfo = TableStyleInfo(name=style["table_style"], showRowStripes=True, showColumnStripes=False)
    sheet.add_table(table)
    sheet.print_options.horizontalCentered = True
    sheet.print_title_rows = "1:1"


def link(cell, target: str, style: dict) -> None:
    cell.hyperlink = "#" + target
    cell.font = Font(name=style["font"], size=style["font_size"], color=style["link_color"], underline="single")


def build_workbook(records: list[dict], config: dict, output: Path, warnings: list[str]) -> dict:
    style = config["style"]
    image_config = config["images"]
    columns = config["columns"]
    workbook = Workbook()
    main = workbook.active
    main.title = "Historique"
    context = workbook.create_sheet("Contexte")
    gallery = workbook.create_sheet("Images")
    info = workbook.create_sheet("Configuration")
    main.append([c["label"] for c in columns])
    context_columns = [
        {"key": "id", "label": "N°", "width": 7},
        {"key": "prompt", "label": "Demande intégrale", "width": 80},
        {"key": "active_file", "label": "Fichier actif", "width": 38},
        {"key": "open_tabs", "label": "Onglets ouverts", "width": 50},
        {"key": "selection", "label": "Sélection dans l'éditeur", "width": 90},
        {"key": "files", "label": "Fichiers joints ou mentionnés", "width": 55},
        {"key": "question", "label": "Question de clarification", "width": 70},
        {"key": "source", "label": "Source de la session", "width": 65},
    ]
    context.append([c["label"] for c in context_columns])
    gallery_columns = [
        {"label": "N° du prompt", "width": 14},
        {"label": "Image", "width": 9},
        {"label": "Capture intégrée", "width": max(60, (image_config["gallery_width_px"] + 18) / 7)},
    ]
    gallery.append([c["label"] for c in gallery_columns])
    gallery_rows = {}
    picture_count = 0
    preview_count = 0
    for row_index, record in enumerate(records, 2):
        values = dict(record)
        source = record["source"]
        values["source"] = f".codex/{source[0]}, ligne {source[1]}" if source else ""
        context.append([values[c["key"]] for c in context_columns])
        context.row_dimensions[row_index].height = style["context_row_height"]
        limit = config["text"]["max_prompt_chars"]
        if limit and len(values["prompt"]) > limit:
            values["prompt"] = values["prompt"][:limit].rstrip() + "…\n\nTexte intégral : onglet Contexte."
        expected = record["expected_images"]
        values["images"] = f"{expected} image(s) — ouvrir la galerie" if expected else ""
        if not image_config["enabled"] and expected:
            values["images"] = f"{expected} image(s) — export désactivé dans la configuration"
        main.append([values[c["key"]] for c in columns])
        lines = max((
            sum(max(1, math.ceil(len(line) / max(1, c["width"] - 3))) for line in str(values[c["key"]]).splitlines())
            for c in columns if c["key"] != "images"
        ), default=1)
        main.row_dimensions[row_index].height = min(style["row_max_height"], max(style["row_min_height"], lines * 14 + 12))
        for image_number, content in enumerate(record["image_data"], 1):
            gallery_row = gallery.max_row + 1
            gallery_rows.setdefault(record["id"], gallery_row)
            gallery.append([record["id"], image_number, ""])
            height = add_picture(gallery, content, 3, gallery_row, image_config["gallery_width_px"], image_config["gallery_height_px"], 8)
            gallery.row_dimensions[gallery_row].height = (height + 20) * 0.75
            picture_count += 1
        image_column = next((i for i, c in enumerate(columns, 1) if c["key"] == "images"), None)
        if image_column and record["image_data"] and image_config["preview_in_table"]:
            # The gallery contains every attachment; the table shows its first image.
            width = min(image_config["preview_width_px"], int(columns[image_column - 1]["width"] * 7 - 16))
            height = add_picture(main, record["image_data"][0], image_column, row_index, width, image_config["preview_height_px"], 30)
            main.row_dimensions[row_index].height = max(main.row_dimensions[row_index].height, (height + 40) * 0.75)
            preview_count += 1

    format_table(main, columns, "Prompts", style)
    format_table(context, context_columns, "Contextes", style)
    if gallery.max_row > 1:
        format_table(gallery, gallery_columns, "Captures", style)
    else:
        gallery.append(["Aucune image exportée."])
        gallery.column_dimensions["A"].width = 45
    for row_index, record in enumerate(records, 2):
        link(context.cell(row_index, 1), f"'Historique'!A{row_index}", style)
        for col_index, col in enumerate(columns, 1):
            cell = main.cell(row_index, col_index)
            if col["key"] in {"id", "prompt"}:
                link(cell, f"'Contexte'!A{row_index}", style)
            elif col["key"] == "images" and record["id"] in gallery_rows:
                link(cell, f"'Images'!C{gallery_rows[record['id']]}", style)
            elif col["key"] == "summary" and record["summary"].startswith(NO_SUMMARY):
                cell.fill = PatternFill("solid", fgColor=style["warning_background"])
    main_rows = {record["id"]: index for index, record in enumerate(records, 2)}
    for row in gallery.iter_rows(min_row=2):
        if row[0].value in main_rows:
            link(row[0], f"'Historique'!A{main_rows[row[0].value]}", style)

    report = {
        "prompts": len(records), "excluded_ids": config["filters"]["exclude_ids"],
        "images": picture_count, "previews": preview_count,
        "without_distinct_summary": sum(r["summary"].startswith(NO_SUMMARY) for r in records),
        "warnings": warnings,
    }
    info.append(["Information", "Valeur"])
    info.append(["Dates", "Heure de Paris ; dates du CSV source."])
    info.append(["Résumés", "Bilans historiques de Codex ; aucun nouvel audit des changements n'est effectué par ce script."])
    info.append(["Demandes longues", "Les textes complets et les sélections IDE restent disponibles dans l'onglet Contexte."])
    info.append(["Images", "Aperçu de la première image dans le tableau ; toutes les pièces jointes dans Images."])
    info.append(["Configuration éditable", "table_config.json ; les chemins relatifs sont résolus depuis ce fichier."])
    for key, value in report.items():
        info.append([key, json.dumps(value, ensure_ascii=False)])
    for key, value in config.items():
        info.append([f"config.{key}", json.dumps(value, ensure_ascii=False, indent=2)])
    format_table(info, [{"width": 28}, {"width": 105}], "Reglages", style)
    for index in range(2, info.max_row + 1):
        info.row_dimensions[index].height = 65
    workbook.active = 0
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.parent.resolve() == Path(__file__).with_name("history").resolve():
        (output.parent / ".gdignore").touch()
    workbook.save(output)
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path, default=Path(__file__).with_name("table_config.json"))
    parser.add_argument("--output", type=Path, help="Autre fichier XLSX, relatif au dossier courant.")
    parser.add_argument("--no-images", action="store_true", help="Ne pas consulter les sessions ni intégrer d'images.")
    args = parser.parse_args()
    config_path = args.config.resolve()
    config = json.loads(config_path.read_text(encoding="utf-8-sig"))
    columns = config["columns"]
    if not columns or len({c["key"] for c in columns}) != len(columns):
        raise ValueError("Les colonnes doivent être présentes et sans doublon.")
    if any(c["key"] not in KNOWN_COLUMNS or c["width"] < 5 for c in columns):
        raise ValueError("Clé de colonne inconnue ou largeur inférieure à 5.")
    if config["filters"]["order"] not in {"ascending", "descending"}:
        raise ValueError("filters.order doit être ascending ou descending.")
    if not 0 < config["style"]["row_min_height"] <= config["style"]["row_max_height"] <= 409:
        raise ValueError("Les hauteurs de ligne doivent être comprises entre 1 et 409 points.")
    for key in ["preview_width_px", "preview_height_px", "gallery_width_px", "gallery_height_px"]:
        if config["images"][key] <= 0:
            raise ValueError(f"images.{key} doit être positif.")
    if config["images"]["gallery_height_px"] > 520 or config["images"]["preview_height_px"] > 500:
        raise ValueError("La hauteur d'image dépasse la hauteur maximale d'une ligne Excel.")
    if args.no_images:
        config["images"]["enabled"] = False
    records = load_records(config, config_path.parent)
    warnings = recover_images(records, config, config_path.parent)
    output = args.output.resolve() if args.output else path_from_config(config["output"], config_path.parent)
    if output.suffix.lower() != ".xlsx":
        raise ValueError("Le fichier de sortie doit avoir l'extension .xlsx.")
    report = build_workbook(records, config, output, warnings)
    output.with_suffix(".rapport.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Créé : {output}")
    print(f"{report['prompts']} prompts ; {report['images']} images ; {report['previews']} aperçus.")
    for warning in warnings:
        print("Attention : " + warning, file=sys.stderr)


if __name__ == "__main__":
    try:
        main()
    except (ValueError, OSError, KeyError, csv.Error) as error:
        raise SystemExit(f"Erreur : {error}") from error
